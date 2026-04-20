import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "..", "outputs")


# ── Main export function ───────────────────────────────────────
def export_documents(mom_json: dict, session_id: str) -> tuple:
    os.makedirs(OUTPUTS_DIR, exist_ok=True)

    # Generate clean filename from session title
    raw_title = (
        mom_json.get("session_title") or
        mom_json.get("title") or
        "Class_Notes"
    )
    # Clean filename — remove special chars, replace spaces with underscores
    import re
    clean_name = re.sub(r'[^\w\s-]', '', raw_title)
    clean_name = re.sub(r'\s+', '_', clean_name.strip())
    clean_name = clean_name[:60]  # max 60 chars

    # Add session_id suffix to avoid collisions
    filename  = f"{clean_name}_{session_id[:8]}"
    docx_path = os.path.join(OUTPUTS_DIR, f"{filename}.docx")

    _generate_docx(mom_json, docx_path)
    return None, f"/outputs/{filename}.docx"


def delete_documents(session_id: str):
    for ext in ["pdf", "docx"]:
        path = os.path.join(OUTPUTS_DIR, f"{session_id}.{ext}")
        if os.path.exists(path):
            os.remove(path)


# ── Helper: check if a field has real content ──────────────────
def _has_content(value) -> bool:
    """Returns True only if value has actual content — not null/empty."""
    if value is None:
        return False
    if isinstance(value, list):
        return len(value) > 0 and any(
            str(v).strip() for v in value if v is not None
        )
    if isinstance(value, str):
        return len(value.strip()) > 0
    if isinstance(value, dict):
        return any(_has_content(v) for v in value.values())
    return bool(value)


def _ensure_list(value) -> list:
    """Helper to ensure we are iterating over a list, not a string."""
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        return [value] if value.strip() else []
    return [str(value)]


# ══ DOCX Generation ════════════════════════════════════════════
def _generate_docx(notes: dict, path: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Pt(50)
        section.bottom_margin = Pt(50)
        section.left_margin   = Pt(70)
        section.right_margin  = Pt(70)

    # ── Title ──────────────────────────────────────────────────
    title = notes.get("session_title") or notes.get("title") or "Class Session Notes"
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold           = True
    run.font.size      = Pt(18)
    run.font.color.rgb = RGBColor(26, 26, 40)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run("ONLINE CLASS SESSION NOTES")
    sub_run.font.size      = Pt(10)
    sub_run.font.color.rgb = RGBColor(83, 74, 183)
    sub_run.bold           = True

    doc.add_paragraph()

    # ── 1. Session Details ─────────────────────────────────────
    # Only show rows that have content
    session_fields = [
        ("Course Name",      notes.get("course_name")),
        ("Subject / Topic",  notes.get("subject_topic")),
        ("Date",             notes.get("date")),
        ("Time",             notes.get("time")),
        ("Platform",         notes.get("platform", "Google Meet")),
        ("Instructor",       notes.get("instructor_name")),
        ("Prepared By",      notes.get("prepared_by", "MoM Generator")),
    ]

    active_fields = [(k, v) for k, v in session_fields if _has_content(v)]

    if active_fields:
        _docx_heading(doc, "1. Session Details")
        table = doc.add_table(rows=len(active_fields), cols=2)
        table.style = "Table Grid"
        for i, (label, value) in enumerate(active_fields):
            row = table.rows[i]
            label_run = row.cells[0].paragraphs[0].add_run(label)
            label_run.bold      = True
            label_run.font.size = Pt(10)
            value_run = row.cells[1].paragraphs[0].add_run(str(value))
            value_run.font.size = Pt(10)
        doc.add_paragraph()

    # ── 2. Session Overview ────────────────────────────────────
    session_overview = _ensure_list(notes.get("session_overview"))
    if session_overview:
        _docx_heading(doc, "2. Session Overview")
        for item in session_overview:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 3. Learning Objectives ─────────────────────────────────
    learning_objectives = _ensure_list(notes.get("learning_objectives"))
    if learning_objectives:
        _docx_heading(doc, "3. Learning Objectives")
        for item in learning_objectives:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 4. Topics Covered ──────────────────────────────────────
    if _has_content(notes.get("topics_covered")):
        _docx_heading(doc, "4. Topics Covered")
        for i, topic in enumerate(notes["topics_covered"], 1):
            if not isinstance(topic, dict):
                continue
            name = topic.get("name", f"Topic {i}")
            # Sub-heading for each topic
            sub = doc.add_paragraph()
            sub_run = sub.add_run(f"4.{i}  {name}")
            sub_run.bold           = True
            sub_run.font.size      = Pt(11)
            sub_run.font.color.rgb = RGBColor(30, 30, 60)

            if _has_content(topic.get("explanation")):
                doc.add_paragraph(f"Explanation: {topic['explanation']}")

            key_points = _ensure_list(topic.get("key_points"))
            if key_points:
                kp_para = doc.add_paragraph()
                kp_para.add_run("Key Points:").bold = True
                for kp in key_points:
                    if _has_content(kp):
                        _docx_bullet(doc, kp)

            examples = _ensure_list(topic.get("examples"))
            if examples:
                ex_para = doc.add_paragraph()
                ex_para.add_run("Examples:").bold = True
                for ex in examples:
                    if _has_content(ex):
                        _docx_bullet(doc, ex)

            if _has_content(topic.get("important_notes")):
                note_para = doc.add_paragraph()
                note_run  = note_para.add_run(f"Important: {topic['important_notes']}")
                note_run.font.color.rgb = RGBColor(180, 50, 50)
                note_run.font.size      = Pt(10)

            doc.add_paragraph()

    # ── 5. Detailed Concepts ───────────────────────────────────
    if _has_content(notes.get("concepts")):
        _docx_heading(doc, "5. Detailed Explanation (Concept Notes)")
        for concept in notes["concepts"]:
            if not isinstance(concept, dict):
                continue
            name = concept.get("name", "Concept")
            c_para = doc.add_paragraph()
            c_run  = c_para.add_run(f"● {name}")
            c_run.bold      = True
            c_run.font.size = Pt(11)

            if _has_content(concept.get("definition")):
                _docx_indented(doc, f"Definition: {concept['definition']}")
            if _has_content(concept.get("explanation")):
                _docx_indented(doc, f"Explanation: {concept['explanation']}")
            if _has_content(concept.get("real_example")):
                _docx_indented(doc, f"Example: {concept['real_example']}")

            doc.add_paragraph()

    # ── 6. Examples / Problems Solved ─────────────────────────
    if _has_content(notes.get("examples")):
        _docx_heading(doc, "6. Examples / Problems Solved")
        for i, ex in enumerate(notes["examples"], 1):
            if not isinstance(ex, dict):
                continue
            ex_title = doc.add_paragraph()
            ex_run   = ex_title.add_run(f"Example {i}:")
            ex_run.bold      = True
            ex_run.font.size = Pt(10)

            if _has_content(ex.get("question")):
                _docx_indented(doc, f"Question: {ex['question']}")
            if _has_content(ex.get("solution_steps")):
                _docx_indented(doc, f"Solution: {ex['solution_steps']}")
            if _has_content(ex.get("final_answer")):
                ans_para = doc.add_paragraph()
                ans_run  = ans_para.add_run(f"    Answer: {ex['final_answer']}")
                ans_run.bold            = True
                ans_run.font.size       = Pt(10)
                ans_run.font.color.rgb  = RGBColor(15, 110, 86)

            doc.add_paragraph()

    # ── 7. Key Takeaways ──────────────────────────────────────
    key_takeaways = _ensure_list(notes.get("key_takeaways"))
    if key_takeaways:
        _docx_heading(doc, "7. Key Takeaways ⭐")
        for item in key_takeaways:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 8. Formulas / Definitions ─────────────────────────────
    formulas_definitions = _ensure_list(notes.get("formulas_definitions"))
    if formulas_definitions:
        _docx_heading(doc, "8. Important Formulas / Definitions")
        for item in formulas_definitions:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 9. Questions & Answers ────────────────────────────────
    if _has_content(notes.get("questions_answers")):
        _docx_heading(doc, "9. Questions & Answers")
        for i, qa in enumerate(notes["questions_answers"], 1):
            if not isinstance(qa, dict):
                continue
            if _has_content(qa.get("question")):
                q_para = doc.add_paragraph()
                q_run  = q_para.add_run(f"Q{i}: {qa['question']}")
                q_run.bold      = True
                q_run.font.size = Pt(10)
            if _has_content(qa.get("answer")):
                _docx_indented(doc, f"A: {qa['answer']}")
        doc.add_paragraph()

    # ── 10. Assignments ───────────────────────────────────────
    assignments = _ensure_list(notes.get("assignments"))
    if assignments:
        _docx_heading(doc, "10. Assignments / Practice Work")
        for item in assignments:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 11. Study Resources ───────────────────────────────────
    study_resources = _ensure_list(notes.get("study_resources"))
    if study_resources:
        _docx_heading(doc, "11. Study Resources")
        for item in study_resources:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 12. Additional Notes ──────────────────────────────────
    additional_notes = _ensure_list(notes.get("additional_notes"))
    if additional_notes:
        _docx_heading(doc, "12. Additional Notes")
        for item in additional_notes:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── 13. Revision Summary ──────────────────────────────────
    revision_summary = _ensure_list(notes.get("revision_summary"))
    if revision_summary:
        _docx_heading(doc, "13. Revision Summary (1-Minute Review)")
        for item in revision_summary:
            if _has_content(item):
                _docx_bullet(doc, item)
        doc.add_paragraph()

    # ── Footer ────────────────────────────────────────────────
    footer      = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run  = footer_para.add_run("Generated by NoteCraft Generator (NCG)")
    footer_run.font.size      = Pt(8)
    footer_run.font.color.rgb = RGBColor(150, 150, 150)

    doc.save(path)


# ── Section heading with purple underline ──────────────────────
def _docx_heading(doc: Document, title: str):
    para = doc.add_paragraph()
    run  = para.add_run(title)
    run.bold           = True
    run.font.size      = Pt(11)
    run.font.color.rgb = RGBColor(83, 74, 183)
    pPr    = para._p.get_or_add_pPr()
    pBdr   = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "534AB7")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Bullet point ───────────────────────────────────────────────
def _docx_bullet(doc: Document, text: str):
    para = doc.add_paragraph(style="List Bullet")
    run  = para.add_run(str(text))
    run.font.size = Pt(10)


# ── Indented text (for sub-items) ─────────────────────────────
def _docx_indented(doc: Document, text: str):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(20)
    run  = para.add_run(str(text))
    run.font.size = Pt(10)